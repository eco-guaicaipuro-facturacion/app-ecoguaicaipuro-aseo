import streamlit as st
import pandas as pd
from datetime import datetime, date
import os
import io
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from db_supabase import get_facturacion, get_suscriptores

st.set_page_config(page_title="Eco Guaicaipuro - Reportes Cloud", page_icon="📊", layout="wide")

LOGO_ECO = "Eco_Guaicaipuro.png"
LOGO_GUAICAIPURO = "Guaicaipuro_Ciudad_Capital.png"

st.markdown("""
    <style>
    .main { background-color: #f4f9f4; }
    .header-box {
        background: linear-gradient(135deg, #1b4332 0%, #2d6a4f 100%);
        padding: 20px; border-radius: 12px; color: white; margin-bottom: 20px;
        display: flex; justify-content: space-between; align-items: center;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
    }
    .header-content { flex-grow: 1; }
    .header-icon { font-size: 45px; margin-left: 20px; text-align: right; }
    .filter-box {
        background-color: #ffffff;
        padding: 15px;
        border-radius: 10px;
        border: 2px solid #52b788;
        margin-bottom: 20px;
        box-shadow: 0 2px 4px rgba(0,0,0,0.05);
    }
    </style>
""", unsafe_allow_html=True)

with st.sidebar:
    st.markdown("### 🏛️ Identidad Institucional")
    if os.path.exists(LOGO_ECO): st.image(LOGO_ECO)
    if os.path.exists(LOGO_GUAICAIPURO): st.image(LOGO_GUAICAIPURO)
    st.markdown("---")
    if st.button("🏠 Ir al Módulo Principal", use_container_width=True):
        st.switch_page("app.py")
    if st.button("📄 Ir a Declaración y Pago", use_container_width=True):
        st.switch_page("pages/Declaracion.py")

st.markdown("""
    <div class="header-box">
        <div class="header-content">
            <h1 style="margin:0; font-size: 24px;">📊 ECO GUAICAIPURO - Reportes y Auditoría Cloud</h1>
            <p style="margin:5px 0 0 0; font-size: 13px; opacity: 0.9;">Memoria Histórica en Tiempo Real (Supabase)</p>
        </div>
        <div class="header-icon">📈📋</div>
    </div>
""", unsafe_allow_html=True)

p1, p2, p3, p4, p5 = st.tabs([
    "📑 Facturación", 
    "💰 Recaudación Diario", 
    "🚨 Morosidad y Estado de Cuenta", 
    "🗺️ Recaudación por Dirección", 
    "🔍 Arqueo y Auditoría Bancaria"
])

def cargar_reporte_facturacion():
    data = get_facturacion()
    if data:
        df_rep = pd.DataFrame(data)
        df_rep.columns = df_rep.columns.str.strip()
        rename_map = {
            "no": "NO", "fecha": "FECHA", "codigo_planilla": "CODIGO PLANILLA",
            "nombre_comercial": "NOMBRE COMERCIAL", "periodo_pagado": "PERIODO PAGADO",
            "total_pagado": "TOTAL PAGADO", "direccion": "DIRECCIÓN"
        }
        df_rep = df_rep.rename(columns=rename_map)
        df_rep["FECHA_DT"] = pd.to_datetime(df_rep["FECHA"], format='%d-%m-%Y', errors='coerce')
        if df_rep["FECHA_DT"].isna().all():
            df_rep["FECHA_DT"] = pd.to_datetime(df_rep["FECHA"], errors='coerce')
        return df_rep
    return pd.DataFrame(columns=["NO", "FECHA", "CODIGO PLANILLA", "NOMBRE COMERCIAL", "PERIODO PAGADO", "TOTAL PAGADO", "DIRECCIÓN"])

def cargar_suscriptores():
    data = get_suscriptores()
    if data:
        df_s = pd.DataFrame(data)
        df_s.columns = df_s.columns.str.strip()
        return df_s
    return pd.DataFrame()

def limpiar_monto(val):
    try:
        v = str(val).replace("BS.", "").replace("Bs.", "").strip()
        return float(v.replace(".", "").replace(",", "."))
    except:
        return 0.0

def aplicar_formato_excel_avanzado(output_io, color_hex="1B4332"):
    wb = load_workbook(output_io)
    ws = wb.active
    fill_color = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
    font_header = Font(name="Helvetica", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Helvetica", size=10)
    thin_border = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'), top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))
    
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill_color
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = font_body
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 5, 16)

    io_final = io.BytesIO()
    wb.save(io_final)
    io_final.seek(0)
    return io_final.getvalue()

def renderizar_filtro_fechas(key_sufijo):
    st.markdown('<div class="filter-box">', unsafe_allow_html=True)
    st.markdown("#### 📅 Filtro Opcional por Fechas")
    tipo_filtro = st.radio("Modo de Visualización:", ["Ver Histórico Completo", "Fecha Específica", "Rango de Fechas"], horizontal=True, key=f"tipo_f_{key_sufijo}")
    f_inicio, f_fin = None, None
    if tipo_filtro == "Fecha Específica":
        f_esp = st.date_input("Seleccione la Fecha:", value=date.today(), key=f"f_esp_{key_sufijo}")
        f_inicio, f_fin = pd.to_datetime(f_esp), pd.to_datetime(f_esp)
    elif tipo_filtro == "Rango de Fechas":
        col_fa, col_fb = st.columns(2)
        with col_fa:
            f_inicio = pd.to_datetime(st.date_input("Desde:", value=date.today(), key=f"f_ini_{key_sufijo}"))
        with col_fb:
            f_fin = pd.to_datetime(st.date_input("Hasta:", value=date.today(), key=f"f_fin_{key_sufijo}"))
    st.markdown('</div>', unsafe_allow_html=True)
    return tipo_filtro, f_inicio, f_fin

df_reporte = cargar_reporte_facturacion()
df_suscriptores = cargar_suscriptores()

with p1:
    st.markdown("### 📑 REPORTE DE FACTURACIÓN CONSECUTIVA (CLOUD)")
    if df_reporte.empty:
        st.warning("⚠️ No hay registros de facturación en la nube todavía.")
    else:
        tipo_f, fini, ffin = renderizar_filtro_fechas("p1")
        df_filtrado = df_reporte.copy()
        if tipo_f != "Ver Histórico Completo" and fini is not None and ffin is not None:
            df_filtrado = df_filtrado[(df_filtrado["FECHA_DT"] >= fini) & (df_filtrado["FECHA_DT"] <= ffin)]
        
        df_visual_p1 = df_filtrado.drop(columns=["FECHA_DT", "id", "creado_en"], errors="ignore")
        st.dataframe(df_visual_p1, use_container_width=True, hide_index=True)

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_visual_p1.to_excel(writer, index=False, sheet_name='Facturacion')
        output.seek(0)
        bytes_archivo = aplicar_formato_excel_avanzado(output, color_hex="1B4332")
        st.download_button("📥 Descargar Reporte de Facturación en Excel", data=bytes_archivo, file_name="Reporte_Facturacion_Cloud.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, type="primary")

with p2:
    st.markdown("### 💰 REPORTE DE RECAUDACIÓN DIARIO (CLOUD)")
    if df_reporte.empty:
        st.warning("⚠️ No hay datos para calcular la recaudación.")
    else:
        df_rec = df_reporte.copy()
        df_rec["MONTO_NUM"] = df_rec["TOTAL PAGADO"].apply(limpiar_monto)
        df_agrupado = df_rec.groupby("FECHA").agg(RECAUDACION_BS=("MONTO_NUM", "sum")).reset_index()
        tasa_ref = st.session_state.get("tasa_mmv_val", 1.0)
        if tasa_ref <= 0: tasa_ref = 1.0

        filas_rec = []
        for i, r in df_agrupado.iterrows():
            bs = r["RECAUDACION_BS"]
            dolar = bs / tasa_ref if tasa_ref > 0 else 0.0
            filas_rec.append({
                "No.": str(i + 1).zfill(3),
                "FECHA": r["FECHA"],
                "RECAUDACIÓN EN BS.": f"Bs. {bs:,.2f}",
                "TASA DEL DÍA": f"{tasa_ref:,.2f}",
                "RECAUDACIÓN EN DÓLARES": f"$ {dolar:,.2f}"
            })
        df_rec_vis = pd.DataFrame(filas_rec)
        st.dataframe(df_rec_vis, use_container_width=True, hide_index=True)

with p3:
    st.markdown("### 🚨 REPORTE DE MOROSIDAD (CLOUD)")
    if df_suscriptores.empty:
        st.warning("⚠️ No se encontró la base de datos de suscriptores en la nube.")
    else:
        col_susc = "suscriptor" if "suscriptor" in df_suscriptores.columns else df_suscriptores.columns[1]
        col_dir = "direccion" if "direccion" in df_suscriptores.columns else df_suscriptores.columns[0]
        
        pagados_por_suscriptor = {}
        if not df_reporte.empty:
            for _, row in df_reporte.iterrows():
                nombre = str(row.get("NOMBRE COMERCIAL", "")).upper()
                periodo = str(row.get("PERIODO PAGADO", ""))
                if nombre not in pagados_por_suscriptor:
                    pagados_por_suscriptor[nombre] = []
                pagados_por_suscriptor[nombre].append(periodo)

        meses_cobranza = ["AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
        lista_morosidad = []
        for _, row in df_suscriptores.iterrows():
            nombre_comercial = str(row.get(col_susc, "")).upper()
            dir_val = str(row.get(col_dir, ""))
            pagos_registrados = " ".join(pagados_por_suscriptor.get(nombre_comercial, []))
            meses_faltantes = [m for m in meses_cobranza if m not in pagos_registrados]
            estatus = "SOLVENTE" if len(meses_faltantes) == 0 else f"MOROSO ({len(meses_faltantes)} meses pendientes)"
            lista_morosidad.append({
                "DIRECCIÓN": dir_val,
                "SUSCRIPTOR": nombre_comercial,
                "ESTATUS GENERAL": estatus,
                "MESES PENDIENTES": ", ".join(meses_faltantes) if meses_faltantes else "NINGUNO"
            })
        st.dataframe(pd.DataFrame(lista_morosidad), use_container_width=True, hide_index=True)

with p4:
    st.markdown("### 🗺️ RECAUDACIÓN POR DIRECCIÓN (CLOUD)")
    if df_reporte.empty:
        st.warning("⚠️ No hay registros de facturación.")
    else:
        df_zona = df_reporte.copy()
        df_zona["MONTO_NUM"] = df_zona["TOTAL PAGADO"].apply(limpiar_monto)
        agrupado_zona = df_zona.groupby("DIRECCIÓN").agg(
            TOTAL_PLANILLAS=("CODIGO PLANILLA", "count"),
            RECAUDACION_TOTAL_BS=("MONTO_NUM", "sum")
        ).reset_index()
        agrupado_zona.columns = ["DIRECCIÓN", "PLANILLAS EMITIDAS", "TOTAL RECAUDADO (BS.)"]
        st.dataframe(agrupado_zona, use_container_width=True, hide_index=True)

with p5:
    st.markdown("### 🔍 ARQUEO BANCARIO Y AUDITORÍA (CLOUD)")
    if df_reporte.empty:
        st.warning("⚠️ No hay datos.")
    else:
        total_recaudado = df_reporte["TOTAL PAGADO"].apply(limpiar_monto).sum()
        total_planillas = len(df_reporte)
        col_a1, col_a2 = st.columns(2)
        col_a1.metric("Total Planillas Emitidas", f"{total_planillas} planillas")
        col_a2.metric("Recaudación Acumulada", f"Bs. {total_recaudado:,.2f}")
